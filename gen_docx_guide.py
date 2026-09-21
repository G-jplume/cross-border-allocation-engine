import os
import re
import sys
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement

BASE = r"c:\Users\Administrator\AppData\Roaming\TRAE SOLO CN\ModularData\ai-agent\work-mode-projects\6aa8bbd5fad07774fb71cfd5"
OUTPUT = os.path.join(BASE, "\u5206\u4ed3\u5360\u6bd4\u8ba1\u7b97\u5f15\u64ce_\u4f7f\u7528\u8bf4\u660e.docx")


def add_toc(doc, guide_md):
    """从 guide.md 解析标题，生成静态目录（无需右键更新）。"""
    p = doc.add_paragraph()
    run = p.add_run("\u76ee\u5f55")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x1a, 0x5c, 0xb0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for line in guide_md.split('\n'):
        line = line.rstrip()
        if line.startswith('## ') and not line.startswith('### '):
            text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0)
            r = p.add_run(text)
            r.font.size = Pt(11)
            r.bold = True
        elif line.startswith('### '):
            text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            r = p.add_run(text)
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def add_page_number_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = p.add_run("\u7b2c ")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run2 = p.add_run()
    run2._r.append(fldChar1)
    run2._r.append(instrText)
    run2._r.append(fldChar2)
    run2.font.size = Pt(9)
    run2.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    run3 = p.add_run(" \u9875 / \u5171 ")
    run3.font.size = Pt(9)
    run3.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'begin')
    instrText2 = OxmlElement('w:instrText')
    instrText2.set(qn('xml:space'), 'preserve')
    instrText2.text = 'NUMPAGES'
    fldChar4 = OxmlElement('w:fldChar')
    fldChar4.set(qn('w:fldCharType'), 'end')
    run4 = p.add_run()
    run4._r.append(fldChar3)
    run4._r.append(instrText2)
    run4._r.append(fldChar4)
    run4.font.size = Pt(9)
    run4.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    run5 = p.add_run(" \u9875")
    run5.font.size = Pt(9)
    run5.font.color.rgb = RGBColor(0x99, 0x99, 0x99)


def add_header(section, text):
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(text)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.autofit = False
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        if col_widths:
            cell.width = Cm(col_widths[i])
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = str(val)
            if col_widths:
                cell.width = Cm(col_widths[ci])
    doc.add_paragraph()


def add_meta_box(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    today = datetime.now().strftime("%Y-%m-%d")
    run = p.add_run(f'\u751f\u6210\u65e5\u671f\uff1a{today}  |  \u968f\u4ee3\u7801\u540c\u6b65\u66f4\u65b0')
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)


def get_guide_md():
    """Read guide.md as the single source of truth."""
    path = os.path.join(BASE, "guide.md")
    with open(path, encoding="utf-8") as f:
        return f.read()


def parse_markdown_table(line_iter):
    """Parse a markdown table from lines, return (headers, rows)."""
    headers = []
    rows = []
    for line in line_iter:
        line = line.strip()
        if not line.startswith('|'):
            break
        cells = [c.strip() for c in line.split('|')[1:-1]]
        if all(set(c) <= set('-: ') for c in cells):
            continue
        if not headers:
            headers = cells
        else:
            rows.append(cells)
    return headers, rows


def gen():
    guide_md = get_guide_md()
    print(f"  GUIDE_MD loaded: {len(guide_md)} chars")

    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    section = doc.sections[0]
    add_header(section, "\u5206\u4ed3\u5360\u6bd4\u8ba1\u7b97\u5f15\u64ce \u2014 \u4f7f\u7528\u8bf4\u660e")
    add_page_number_footer(section)

    # Title page
    for _ in range(6):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('\u5206\u4ed3\u5360\u6bd4\u8ba1\u7b97\u5f15\u64ce')
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(0x1a, 0x5c, 0xb0)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run('\u4f7f\u7528\u8bf4\u660e')
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_paragraph()

    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = intro.add_run('\u4e0a\u4f20\u5386\u53f2\u51fa\u5355\u6570\u636e \u2192 \u8c03\u8282\u53c2\u6570 \u2192 \u8ba1\u7b97\u5404\u4ed3\u5e93\u53d1\u8d27\u5360\u6bd4 \u2192 \u5bfc\u51fa\u7ed3\u679c')
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    intro2 = doc.add_paragraph()
    intro2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = intro2.add_run('\u56db\u4ed3\u5360\u6bd4\u5408\u8ba1\u6052\u4e3a 100%  |  \u6700\u5927\u4f59\u989d\u6cd5\u4fdd\u8bc1\u843d\u8d27\u91cf\u6574\u6570\u7cbe\u786e  |  \u81ea\u52a8\u4f53\u68c0\u9632\u9759\u9ed8\u5931\u6548')
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.add_paragraph()
    add_meta_box(doc)

    doc.add_page_break()

    # TOC
    add_toc(doc, guide_md)
    doc.add_page_break()

    # Parse GUIDE_MD markdown and convert to Word
    lines = guide_md.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        # Skip empty lines and horizontal rules
        if not line.strip() or line.strip() == '---':
            i += 1
            continue

        # Headings
        if line.startswith('### '):
            heading_text = line[4:].strip()
            doc.add_heading(heading_text, level=2)
            i += 1
        elif line.startswith('## '):
            heading_text = line[3:].strip()
            doc.add_heading(heading_text, level=1)
            i += 1
        elif line.startswith('# '):
            heading_text = line[2:].strip()
            doc.add_heading(heading_text, level=1)
            i += 1
        # Tables
        elif line.startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i])
                i += 1
            headers, rows = parse_markdown_table(iter(table_lines))
            if headers:
                n_cols = len(headers)
                if n_cols == 3:
                    col_widths = [5, 2, 10]
                elif n_cols == 2:
                    col_widths = [4, 12]
                else:
                    col_widths = [16 / n_cols] * n_cols
                add_table(doc, headers, rows, col_widths)
        # Bullet points
        elif line.startswith('- '):
            text = line[2:].strip()
            # Remove markdown bold
            text = text.replace('**', '')
            doc.add_paragraph(text, style='List Bullet')
            i += 1
        # Numbered lists
        elif re.match(r'^\d+\.\s', line):
            text = re.sub(r'^\d+\.\s', '', line).strip()
            text = text.replace('**', '')
            doc.add_paragraph(text, style='List Number')
            i += 1
        # Regular paragraphs
        else:
            # Collect consecutive non-empty, non-special lines
            para_lines = []
            while i < len(lines):
                l = lines[i].rstrip()
                if not l.strip() or l.startswith('|') or l.startswith('#') or l.startswith('- ') or l.startswith('---') or re.match(r'^\d+\.\s', l):
                    break
                para_lines.append(l)
                i += 1
            if para_lines:
                text = '\n'.join(para_lines)
                # Remove markdown bold markers for display
                text = text.replace('**', '')
                doc.add_paragraph(text)

    doc.save(OUTPUT)
    print(f"  DOCX saved: {OUTPUT}")
    print(f"  DOCX size: {os.path.getsize(OUTPUT)} bytes")


if __name__ == "__main__":
    gen()
